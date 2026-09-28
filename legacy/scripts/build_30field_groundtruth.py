#!/usr/bin/env python3
"""
Build 30-field ground truth for the 8 benchmark datasets.

The existing groundtruth covers 16 fields. This script adds the missing 14 RAI
fields by extracting reference answers from the papers using Claude, then
merging with existing groundtruth. The generated answers should be human-reviewed.

Strategy:
  1. Load existing 16-field groundtruth from groundtruth/parsed_md_filtered/
  2. For each benchmark paper, extract the 14 missing fields using a STRICT
     extraction prompt (different from the pipeline prompt to avoid circularity)
  3. Merge into 30-field groundtruth files
  4. Generate review sheet for human verification

Usage:
    python scripts/build_30field_groundtruth.py                # extract + merge
    python scripts/build_30field_groundtruth.py --merge-only   # just merge existing
    python scripts/build_30field_groundtruth.py --review-sheet  # generate xlsx for review
    python scripts/build_30field_groundtruth.py --stats         # show coverage stats
"""

import sys
import json
import time
import argparse
from pathlib import Path
from collections import defaultdict

sys.path.insert(0, str(Path(__file__).parent.parent))

# ═══════════════════════════════════════════════════════════════════════
# Constants
# ═══════════════════════════════════════════════════════════════════════

BENCHMARK_MAP = {
    "2012.03411v2": "MLS",
    "2009.03300v3": "MMLU",
    "2106.03193v1": "FLORES",
    "2404.00498v2": "CIFAR",
    "1405.0312v3": "MSCOCO",
    "2311.16502v4": "MMMU",
    "1602.07332v1": "Visual Genome",
    "2310.02255v3": "MathVista",
}

RAW_DIR = Path("data/raw")
EXISTING_GT_DIR = Path("groundtruth/parsed_md_filtered")
OUTPUT_DIR = Path("data/groundtruth_30field")
EXTRACTED_DIR = OUTPUT_DIR / "extracted_rai"

ALL_30_FIELDS = [
    "sc:name", "sc:description", "sc:url", "sc:license", "sc:creator", "sc:publisher",
    "sc:datePublished", "sc:inLanguage", "cr:citeAs", "cr:isLiveDataset",
    "rai:dataCollection", "rai:dataCollectionType", "rai:dataCollectionMissingData",
    "rai:dataCollectionRawData", "rai:dataCollectionTimeframe", "rai:dataImputationProtocol",
    "rai:dataManipulationProtocol", "rai:dataPreprocessingProtocol",
    "rai:dataAnnotationProtocol", "rai:dataAnnotationPlatform", "rai:dataAnnotationAnalysis",
    "rai:annotationsPerItem", "rai:annotatorDemographics", "rai:machineAnnotationTools",
    "rai:dataReleaseMaintenancePlan", "rai:personalSensitiveInformation",
    "rai:dataSocialImpact", "rai:dataBiases", "rai:dataLimitations", "rai:dataUseCases",
]

# The 14 fields missing from existing groundtruth
MISSING_FIELDS = [
    "rai:dataCollectionType", "rai:dataCollectionMissingData", "rai:dataCollectionRawData",
    "rai:dataImputationProtocol", "rai:dataManipulationProtocol", "rai:dataPreprocessingProtocol",
    "rai:dataAnnotationProtocol", "rai:dataAnnotationAnalysis", "rai:annotationsPerItem",
    "rai:machineAnnotationTools", "rai:dataReleaseMaintenancePlan",
    "rai:dataSocialImpact", "rai:dataBiases", "rai:dataLimitations",
]

# Strict extraction prompt — intentionally different from the pipeline prompt
# to avoid circularity. Emphasizes quoting the paper directly.
GT_SYSTEM_PROMPT = """You are creating GROUND TRUTH annotations for evaluating an ML metadata extraction system.

Your task: read the paper carefully and extract the CORRECT value for each requested field.
These annotations will be used to judge whether an automated system extracted correctly.

CRITICAL RULES:
1. Quote or closely paraphrase the paper. Do not invent, generalize, or infer.
2. If the paper does not discuss a topic, the correct answer is null.
3. Be thorough — check the entire paper including appendices, ethics statements, limitations sections.
4. For null answers, briefly explain WHY it's null (e.g., "Paper does not discuss maintenance plans").
5. Return valid JSON only."""

GT_USER_PROMPT = """Read this paper and extract ground truth values for these 14 RAI metadata fields.
For each field, provide the CORRECT value based on what the paper explicitly states.

FIELDS TO EXTRACT:
{field_definitions}

Return JSON with this exact structure:
{{
  "rai:dataCollectionType": "value or null",
  "rai:dataCollectionMissingData": "value or null",
  "rai:dataCollectionRawData": "value or null",
  "rai:dataImputationProtocol": "value or null",
  "rai:dataManipulationProtocol": "value or null",
  "rai:dataPreprocessingProtocol": "value or null",
  "rai:dataAnnotationProtocol": "value or null",
  "rai:dataAnnotationAnalysis": "value or null",
  "rai:annotationsPerItem": "value or null",
  "rai:machineAnnotationTools": "value or null",
  "rai:dataReleaseMaintenancePlan": "value or null",
  "rai:dataSocialImpact": "value or null",
  "rai:dataBiases": "value or null",
  "rai:dataLimitations": "value or null"
}}

PAPER TEXT:
{paper_text}"""

FIELD_DEFINITIONS = """
- rai:dataCollectionType: Type of data collection. Choose from: Surveys, Secondary Data analysis, Physical data collection, Direct measurement, Document analysis, Manual Human Curator, Software Collection, Experiments, Web Scraping, Web API, Focus groups, Self-reporting, Customer feedback data, User-generated content data, Passive Data Collection, Others.
- rai:dataCollectionMissingData: How missing data was handled. null if not discussed.
- rai:dataCollectionRawData: Description of the raw/source data before processing.
- rai:dataImputationProtocol: How missing values were imputed. null if not discussed.
- rai:dataManipulationProtocol: Data modifications after preprocessing (augmentation, balancing, etc.). null if not discussed.
- rai:dataPreprocessingProtocol: Steps to make raw data ML-ready (filtering, cleaning, formatting).
- rai:dataAnnotationProtocol: How labels/annotations were created. Include annotator type, instructions, QC. null if no annotation step.
- rai:dataAnnotationAnalysis: How raw annotations were converted to final labels (agreement analysis, adjudication). null if not discussed.
- rai:annotationsPerItem: Number of human annotations per data item. null if not stated.
- rai:machineAnnotationTools: Software tools used for automated annotation. null if none mentioned.
- rai:dataReleaseMaintenancePlan: Versioning, update schedule, maintenance team. null if not discussed (most papers don't discuss this).
- rai:dataSocialImpact: Social implications discussed by the authors. null if not discussed.
- rai:dataBiases: Biases explicitly acknowledged by the authors. null if not discussed.
- rai:dataLimitations: Limitations explicitly stated by the authors. null if not discussed.
"""


def load_existing_groundtruth():
    """Load existing 16-field groundtruth."""
    gt_file = EXISTING_GT_DIR / "all_annotations.json"
    if not gt_file.exists():
        print(f"  Groundtruth file not found: {gt_file}")
        return {}
    with open(gt_file) as f:
        return json.load(f)


def extract_missing_fields(ds_name, pdf_id):
    """Extract the 14 missing RAI fields from a paper using Claude."""
    from metadata.extractor import setup_llm_pipeline, clean_llm_output
    from pdf.reader import extract_text_from_pdf
    from pdf.processor import clean_text
    from config import MAX_PDF_CHARS
    import re

    pdf_path = RAW_DIR / f"{pdf_id}.pdf"
    if not pdf_path.exists():
        print(f"  PDF not found: {pdf_path}")
        return None

    raw_text = extract_text_from_pdf(pdf_path)
    cleaned = clean_text(raw_text)
    if len(cleaned) > MAX_PDF_CHARS:
        cleaned = cleaned[:MAX_PDF_CHARS]

    user_prompt = GT_USER_PROMPT.format(
        field_definitions=FIELD_DEFINITIONS,
        paper_text=cleaned
    )

    model = setup_llm_pipeline("claude-sonnet-4-5")
    output = model.generate(user_prompt, system_prompt=GT_SYSTEM_PROMPT)
    json_str = clean_llm_output(output, user_prompt)

    try:
        return json.loads(json_str)
    except json.JSONDecodeError:
        # Regex fallback
        result = {}
        for m in re.finditer(r'"([^"]+)"\s*:\s*null', json_str):
            result[m.group(1)] = None
        for m in re.finditer(r'"([^"]+)"\s*:\s*"((?:[^"\\]|\\.)*)"', json_str):
            result[m.group(1)] = m.group(2).replace('\\n', ' ').strip()
        return result


def run_extraction():
    """Extract missing fields for all 8 benchmark datasets."""
    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)

    for pdf_id, ds_name in BENCHMARK_MAP.items():
        out_file = EXTRACTED_DIR / f"{ds_name}_rai_gt.json"
        if out_file.exists():
            print(f"  {ds_name}: already extracted, skipping")
            continue

        print(f"  Extracting {ds_name} ({pdf_id}.pdf)...", end=" ", flush=True)
        start = time.time()

        result = extract_missing_fields(ds_name, pdf_id)
        if result:
            with open(out_file, "w", encoding="utf-8") as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            n_filled = sum(1 for v in result.values() if v is not None)
            print(f"{n_filled}/14 fields filled ({time.time()-start:.0f}s)")
        else:
            print("FAILED")


def merge_groundtruth():
    """Merge existing 16-field GT with extracted 14 RAI fields into 30-field GT."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    existing_gt = load_existing_groundtruth()
    benchmarks = list(BENCHMARK_MAP.values())

    merged = {}
    stats = {"datasets": 0, "total_fields": 0, "filled_fields": 0}

    for ds_name in benchmarks:
        # Start with existing annotations
        existing_anns = existing_gt.get(ds_name, [])

        # Load extracted RAI fields
        rai_file = EXTRACTED_DIR / f"{ds_name}_rai_gt.json"
        rai_fields = {}
        if rai_file.exists():
            with open(rai_file) as f:
                rai_fields = json.load(f)

        # Create merged annotations
        # For existing annotations, add the RAI fields to each annotator's entry
        merged_anns = []
        if existing_anns:
            for ann in existing_anns:
                merged_ann = dict(ann)
                # Add RAI fields (same values for all annotators since we have 1 extraction)
                for field in MISSING_FIELDS:
                    if field not in merged_ann or not merged_ann.get(field):
                        merged_ann[field] = rai_fields.get(field)
                merged_anns.append(merged_ann)
        else:
            # No existing GT — create a single annotator entry from RAI extraction
            merged_ann = {"dataset_name": ds_name, "annotator_id": "claude_gt_extraction"}
            for field in ALL_30_FIELDS:
                merged_ann[field] = rai_fields.get(field)
            merged_anns = [merged_ann]

        merged[ds_name] = merged_anns

        # Count coverage
        stats["datasets"] += 1
        all_fields_in_ds = set()
        for ann in merged_anns:
            for f in ALL_30_FIELDS:
                if f in ann and ann[f] is not None and str(ann[f]).strip():
                    all_fields_in_ds.add(f)
        stats["total_fields"] += 30
        stats["filled_fields"] += len(all_fields_in_ds)

    # Save merged groundtruth
    # Per-dataset files
    for ds_name, anns in merged.items():
        ds_file = OUTPUT_DIR / f"{ds_name.lower().replace(' ', '_')}_annotations.json"
        with open(ds_file, "w", encoding="utf-8") as f:
            json.dump(anns, f, indent=2, ensure_ascii=False)

    # Combined file
    with open(OUTPUT_DIR / "all_annotations.json", "w", encoding="utf-8") as f:
        json.dump(merged, f, indent=2, ensure_ascii=False)

    print(f"\n  Merged groundtruth saved to {OUTPUT_DIR}/")
    print(f"  Datasets: {stats['datasets']}")
    print(f"  Field coverage: {stats['filled_fields']}/{stats['total_fields']} "
          f"({stats['filled_fields']/stats['total_fields']*100:.0f}%)")

    return merged


def show_stats():
    """Show coverage statistics for the merged groundtruth."""
    merged_file = OUTPUT_DIR / "all_annotations.json"
    if not merged_file.exists():
        print("No merged groundtruth found. Run without --stats first.")
        return

    with open(merged_file) as f:
        merged = json.load(f)

    print(f"\n{'='*80}")
    print("30-FIELD GROUNDTRUTH COVERAGE")
    print(f"{'='*80}")

    # Per-dataset coverage
    print(f"\n{'Dataset':<20} {'Fields':>7} {'General':>9} {'RAI':>5} {'Missing fields'}")
    print("-" * 80)

    total_filled = 0
    total_possible = 0

    for ds_name in BENCHMARK_MAP.values():
        anns = merged.get(ds_name, [])
        filled = set()
        for ann in anns:
            for f in ALL_30_FIELDS:
                if f in ann and ann[f] is not None and str(ann[f]).strip():
                    filled.add(f)

        general = [f for f in filled if not f.startswith("rai:")]
        rai = [f for f in filled if f.startswith("rai:")]
        missing = [f for f in ALL_30_FIELDS if f not in filled]

        total_filled += len(filled)
        total_possible += 30

        miss_str = ", ".join(missing[:5])
        if len(missing) > 5:
            miss_str += f" (+{len(missing)-5} more)"

        print(f"{ds_name:<20} {len(filled):>4}/30 {len(general):>6}/10 {len(rai):>3}/20 {miss_str}")

    print("-" * 80)
    print(f"{'TOTAL':<20} {total_filled:>4}/{total_possible} ({total_filled/total_possible*100:.0f}%)")

    # Per-field coverage (across datasets)
    print(f"\n{'Field':<40} {'Datasets with value':>20}")
    print("-" * 65)

    for field in ALL_30_FIELDS:
        count = 0
        for ds_name in BENCHMARK_MAP.values():
            anns = merged.get(ds_name, [])
            for ann in anns:
                if field in ann and ann[field] is not None and str(ann[field]).strip():
                    count += 1
                    break
        marker = "" if count == 8 else " ◄" if count == 0 else ""
        print(f"{field:<40} {count:>4}/8{marker}")


def generate_review_sheet():
    """Generate an Excel sheet for human review of the extracted RAI groundtruth."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("openpyxl required. pip install openpyxl")
        return

    merged_file = OUTPUT_DIR / "all_annotations.json"
    if not merged_file.exists():
        print("No merged groundtruth found. Run extraction + merge first.")
        return

    with open(merged_file) as f:
        merged = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "30-Field GT Review"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    new_fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    existing_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
    null_fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    headers = ["#", "Dataset", "Field", "Source", "Value", "Correct? (Y/N)", "Corrected Value", "Notes"]
    widths = [5, 18, 38, 12, 80, 14, 60, 35]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A" + chr(64 + col - 26)].width = w

    ws.freeze_panes = "A2"

    row_num = 2
    for ds_name in BENCHMARK_MAP.values():
        anns = merged.get(ds_name, [])
        if not anns:
            continue

        # Use first annotator's values
        ann = anns[0]

        for field in ALL_30_FIELDS:
            value = ann.get(field)
            is_new = field in MISSING_FIELDS
            source = "AI-extracted" if is_new else "Human GT"

            ws.cell(row=row_num, column=1, value=row_num - 1).border = thin_border
            ws.cell(row=row_num, column=2, value=ds_name).border = thin_border
            cell_field = ws.cell(row=row_num, column=3, value=field)
            cell_field.font = Font(name="Arial", size=10, bold=True)
            cell_field.border = thin_border

            cell_source = ws.cell(row=row_num, column=4, value=source)
            cell_source.fill = new_fill if is_new else existing_fill
            cell_source.border = thin_border

            display_value = str(value)[:500] if value else "[NULL]"
            cell_val = ws.cell(row=row_num, column=5, value=display_value)
            cell_val.alignment = Alignment(wrap_text=True, vertical="top")
            cell_val.border = thin_border
            if value is None:
                cell_val.fill = null_fill

            for col in [6, 7, 8]:
                ws.cell(row=row_num, column=col).border = thin_border

            ws.row_dimensions[row_num].height = 35
            row_num += 1

    out_file = OUTPUT_DIR / "30field_review_sheet.xlsx"
    wb.save(out_file)
    print(f"  Review sheet saved to {out_file}")
    print(f"  Total rows: {row_num - 2} (8 datasets × 30 fields = 240)")
    print(f"  Yellow = AI-extracted (needs review), Green = existing human GT")


def main():
    parser = argparse.ArgumentParser(description="Build 30-field ground truth")
    parser.add_argument("--merge-only", action="store_true",
                        help="Skip extraction, just merge existing files")
    parser.add_argument("--review-sheet", action="store_true",
                        help="Generate Excel review sheet")
    parser.add_argument("--stats", action="store_true",
                        help="Show coverage statistics")
    args = parser.parse_args()

    if args.stats:
        show_stats()
        return

    if args.review_sheet:
        generate_review_sheet()
        return

    print(f"{'='*60}")
    print("BUILD 30-FIELD GROUNDTRUTH")
    print(f"{'='*60}")

    if not args.merge_only:
        print("\nStep 1: Extract missing 14 RAI fields from papers...")
        run_extraction()
    else:
        print("\nStep 1: Skipped (--merge-only)")

    print("\nStep 2: Merge with existing 16-field groundtruth...")
    merge_groundtruth()

    print("\nStep 3: Generate review sheet...")
    generate_review_sheet()

    print("\nStep 4: Coverage statistics...")
    show_stats()

    print(f"\n{'='*60}")
    print("NEXT STEPS:")
    print(f"{'='*60}")
    print("1. Review data/groundtruth_30field/30field_review_sheet.xlsx")
    print("2. Mark each AI-extracted field as Y (correct) or N (needs correction)")
    print("3. Fill in corrections where needed")
    print("4. Run evaluation with: --groundtruth-dir data/groundtruth_30field")


if __name__ == "__main__":
    main()
