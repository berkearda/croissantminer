#!/usr/bin/env python3
"""Aggregate returned re-annotation pilot sheets into an alternative gold
and analyse it (NeurIPS rebuttal, T-088 pilot).

Stages (all read-only w.r.t. existing gold; outputs under data/rebuttal_pilot/):
  1. parse    - read completed sheets from data/rebuttal_pilot/returns/,
                one xlsx per (paper, slot); extract per-field rating,
                failure mode, corrected value, confidence, rater name.
  2. aggregate- majority vote per cell using the same rule as
                scripts/annotations/build_gold_iaa.py: mode rating 1 keeps
                the pre-fill as gold; otherwise the most common corrected
                value among mode voters; no majority -> tie (adjudication).
  3. compare  - field-level agreement between pilot gold and current gold
                (exact match and null-status agreement per field).
  4. score    - re-score systems against pilot gold on the pilot papers:
                Tier 1 rules locally; Tier 2 reads judge scores from
                data/judged/judge_scores_pilot_glm5.parquet where present
                (produced by scripts/judge_pilot_gold.py) and reports
                which cells still need judging.

Usage:
    python scripts/pilot_aggregate_and_score.py parse
    python scripts/pilot_aggregate_and_score.py aggregate
    python scripts/pilot_aggregate_and_score.py compare
    python scripts/pilot_aggregate_and_score.py score
    python scripts/pilot_aggregate_and_score.py all
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import re
import sys
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

PILOT_DIR = ROOT / "data" / "rebuttal_pilot"
RETURNS = PILOT_DIR / "returns"
SELECTION = ROOT / "data" / "rebuttal_pilot_selection_2026-07-27.csv"

spec = importlib.util.spec_from_file_location(
    "btb", ROOT / "scripts" / "figures" / "build_test88_headline_table.py")
btb = importlib.util.module_from_spec(spec)
spec.loader.exec_module(btb)

# canonical prefixed field ids in sheet order (sheets use unprefixed g2 keys)
spec2 = importlib.util.spec_from_file_location(
    "g2", ROOT / "scripts" / "generate_phase2_sheets.py")
g2 = importlib.util.module_from_spec(spec2)
spec2.loader.exec_module(g2)

RATING_RE = re.compile(r"^\s*([123])")


def parse_rating(v):
    if v is None:
        return None
    m = RATING_RE.match(str(v))
    return int(m.group(1)) if m else None


def stage_parse():
    rows = []
    files = sorted(RETURNS.glob("*.xlsx"))
    if not files:
        print(f"no returned sheets in {RETURNS}")
        return pd.DataFrame()
    for f in files:
        m = re.match(r"(\d+)_(.+)_slot([ABC])", f.stem)
        if not m:
            print(f"SKIP unrecognised file name: {f.name}")
            continue
        paper, slot = m.group(2), m.group(3)
        wb = load_workbook(f, data_only=True)
        ws = wb["Annotations"]
        n_rated = 0
        for r in range(2, ws.max_row + 1):
            field = ws.cell(r, 4).value
            if not field:
                continue
            rating = parse_rating(ws.cell(r, 6).value)
            if rating is not None:
                n_rated += 1
            rows.append({
                "paper_id": paper, "slot": slot, "field": str(field).strip(),
                "prefill": ws.cell(r, 5).value,
                "rating": rating,
                "failure_mode": ws.cell(r, 7).value,
                "corrected_value": ws.cell(r, 8).value,
                "confidence": ws.cell(r, 9).value,
                "notes": ws.cell(r, 10).value,
                "source_file": f.name,
            })
        print(f"{f.name}: {n_rated}/30 rated")
    df = pd.DataFrame(rows)
    for c in ("prefill", "failure_mode", "corrected_value", "confidence", "notes"):
        df[c] = df[c].map(lambda v: None if v is None else str(v))
    df.to_parquet(PILOT_DIR / "pilot_ratings.parquet", index=False)
    print(f"parsed {len(files)} sheets -> {len(df)} rows "
          f"-> {PILOT_DIR/'pilot_ratings.parquet'}")
    return df


def aggregate_cell(grp: pd.DataFrame):
    """Mirror build_gold_iaa.py: mode rating; 1 -> prefill is gold; else the
    most common corrected value among mode voters; tie -> adjudication."""
    ratings = [r for r in grp["rating"] if r is not None]
    if len(ratings) < 2:
        return {"gold_value": None, "method": "insufficient_ratings",
                "n_ratings": len(ratings)}
    counts = Counter(ratings)
    top, top_n = counts.most_common(1)[0]
    tied = [v for v, c in counts.items() if c == top_n]
    if len(tied) > 1:
        return {"gold_value": None, "method": "tie_split",
                "n_ratings": len(ratings)}
    if top == 1:
        return {"gold_value": str(grp["prefill"].iloc[0]),
                "method": f"majority_{top_n}of{len(ratings)}_keep",
                "n_ratings": len(ratings)}
    voters = grp[grp["rating"] == top]
    corrs = [str(c).strip() for c in voters["corrected_value"]
             if c is not None and str(c).strip()]
    if not corrs:
        # flagged but nobody wrote a correction: rating 2 keeps prefill
        # with a note, rating 3 nulls the cell (merge_adjudication rules)
        if top == 2:
            return {"gold_value": str(grp["prefill"].iloc[0]),
                    "method": f"majority_{top_n}of{len(ratings)}_partial_keep",
                    "n_ratings": len(ratings)}
        return {"gold_value": "[null - not found in paper]",
                "method": f"majority_{top_n}of{len(ratings)}_nulled",
                "n_ratings": len(ratings)}
    val, _ = Counter(corrs).most_common(1)[0]
    return {"gold_value": val,
            "method": f"majority_{top_n}of{len(ratings)}_corrected",
            "n_ratings": len(ratings)}


FIELD_TO_ID = {}
for fid in pd.read_parquet(ROOT / "data/annotations/gold.parquet").field_id.unique():
    FIELD_TO_ID[fid] = fid
    FIELD_TO_ID[fid.split(":")[-1]] = fid


def stage_aggregate():
    df = pd.read_parquet(PILOT_DIR / "pilot_ratings.parquet")
    out = []
    for (paper, field), grp in df.groupby(["paper_id", "field"]):
        rec = aggregate_cell(grp)
        rec.update(paper_id=paper, field_id=FIELD_TO_ID.get(field, field))
        out.append(rec)
    g = pd.DataFrame(out)
    g.to_parquet(PILOT_DIR / "pilot_gold.parquet", index=False)
    ties = g[g.method == "tie_split"]
    print(f"aggregated {len(g)} cells -> {PILOT_DIR/'pilot_gold.parquet'}")
    print(f"methods: {g.method.value_counts().to_dict()}")
    if len(ties):
        ties.to_csv(PILOT_DIR / "pilot_adjudication_needed.csv", index=False)
        print(f"ADJUDICATION NEEDED: {len(ties)} cells -> "
              f"{PILOT_DIR/'pilot_adjudication_needed.csv'}")
    return g


def stage_compare():
    pilot = pd.read_parquet(PILOT_DIR / "pilot_gold.parquet")
    cur = pd.read_parquet(ROOT / "data/annotations/gold.parquet")
    m = pilot.merge(cur[["paper_id", "field_id", "gold_value"]],
                    on=["paper_id", "field_id"], suffixes=("_pilot", "_cur"))
    m = m[m.gold_value_pilot.notna()]
    pn = m.gold_value_pilot.map(btb.is_null_gold)
    cn = m.gold_value_cur.map(btb.is_null_gold)
    null_agree = (pn == cn)
    exact = (m.gold_value_pilot.fillna("").str.strip().str.lower()
             == m.gold_value_cur.fillna("").str.strip().str.lower())
    print(f"cells compared: {len(m)}")
    print(f"null-status agreement: {null_agree.mean():.3f}")
    print(f"exact value agreement: {exact.mean():.3f} "
          f"(lexical; semantic agreement requires the judge)")
    per_field = pd.DataFrame({
        "field_id": m.field_id, "null_agree": null_agree, "exact": exact,
    }).groupby("field_id").mean().sort_values("null_agree")
    print(per_field.to_string())
    per_field.to_csv(PILOT_DIR / "pilot_vs_current_agreement.csv")
    return m


def stage_score():
    pilot = pd.read_parquet(PILOT_DIR / "pilot_gold.parquet")
    papers = sorted(pilot.paper_id.unique())
    print(f"scoring against pilot gold on {len(papers)} papers")
    judge_path = ROOT / "data/judged/judge_scores_pilot_glm5.parquet"
    judged = pd.read_parquet(judge_path) if judge_path.exists() else None
    from evaluation.field_metrics import score_field
    SYSTEMS = list(btb.STRATEGY_DIRS)
    need_judge = []
    results = {}
    SM = {1: 1.0, 2: 0.5, 3: 0.0}
    for sysid in SYSTEMS:
        sdir = btb.STRATEGY_DIRS[sysid]
        if not sdir.exists():
            continue
        cells = []
        for _, g in pilot[pilot.gold_value.notna()].iterrows():
            path = sdir / f"{g.paper_id}.json"
            if not path.exists():
                continue
            payload = json.load(open(path))
            ext = payload.get("extraction", payload)
            if not isinstance(ext, dict):
                ext = {}
            fid = g.field_id
            cand = ext.get(fid, ext.get(fid.split(":")[-1]))
            cand_empty = cand is None or not str(cand).strip()
            gold_null = btb.is_null_gold(g.gold_value)
            if gold_null:
                if cand_empty:
                    continue
                cells.append((g.paper_id, fid, 0.0))
                continue
            if fid in btb.TIER1:
                r = score_field(cand, g.gold_value, fid)
                if not r["skipped"]:
                    cells.append((g.paper_id, fid, r["score"]))
            else:
                if cand_empty:
                    cells.append((g.paper_id, fid, 0.0))
                elif judged is not None:
                    hit = judged[(judged.system_id == sysid)
                                 & (judged.paper_id == g.paper_id)
                                 & (judged.field_id == fid)]
                    if len(hit):
                        cells.append((g.paper_id, fid,
                                      SM.get(int(hit.score.iloc[0]), 0.0)))
                    else:
                        need_judge.append((sysid, g.paper_id, fid))
                else:
                    need_judge.append((sysid, g.paper_id, fid))
        if cells:
            df = pd.DataFrame(cells, columns=["paper_id", "field_id", "score"])
            results[sysid] = df.score.mean()
    print(f"cells still needing the judge: {len(need_judge)}")
    if need_judge:
        pd.DataFrame(need_judge, columns=["system_id", "paper_id", "field_id"]) \
            .to_csv(PILOT_DIR / "pilot_judge_todo.csv", index=False)
    if results:
        res = pd.Series(results).sort_values(ascending=False)
        print("\nmean cell score vs PILOT gold (judged cells only where needed):")
        print(res.round(3).to_string())
        res.to_csv(PILOT_DIR / "pilot_system_scores.csv")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=["parse", "aggregate", "compare",
                                      "score", "all"])
    args = ap.parse_args()
    RETURNS.mkdir(exist_ok=True)
    stages = (["parse", "aggregate", "compare", "score"]
              if args.stage == "all" else [args.stage])
    for s in stages:
        print(f"\n===== {s} =====")
        globals()[f"stage_{s}"]()


if __name__ == "__main__":
    main()
