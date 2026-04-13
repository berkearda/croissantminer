#!/usr/bin/env python3
"""
Update Phase 2 annotation sheets based on supervisor feedback.

Changes:
  1. Rating scale 1-5 → 1-3 (with descriptive dropdown labels)
  2. Add Field Definitions tab
  3. Add Paper URL column (after Dataset ID)
  4. Update Instructions tab
"""

import re
from pathlib import Path
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).parent.parent
PHASE2_DIR = ROOT / "data" / "annotations" / "phase2"

# ═══════════════════════════════════════════════════════════════════════
# Paper URL mapping (dataset_id → paper URL)
# ═══════════════════════════════════════════════════════════════════════

PAPER_URLS = {
    "CIFAR_30field": "https://arxiv.org/abs/2404.00498",
    "MMLU_30field": "https://arxiv.org/abs/2009.03300",
    "MMMU_30field": "https://arxiv.org/abs/2311.16502",
    "MathVista_30field": "https://arxiv.org/abs/2310.02255",
    "MLS_30field": "https://arxiv.org/abs/2012.03411",
    "FLORES_30field": "https://arxiv.org/abs/2106.03193",
    "MSCOCO_30field": "https://arxiv.org/abs/1405.0312",
    "Visual_Genome_30field": "https://arxiv.org/abs/1602.07332",
    "Rowan_hellaswag": "https://arxiv.org/abs/1905.07830",
    "allenai_ai2_arc": "https://arxiv.org/abs/1803.05457",
    "allenai_openbookqa": "https://arxiv.org/abs/1809.02789",
    "bigcode_bigcodebench": "https://arxiv.org/abs/2406.15877",
    "google_boolq": "https://arxiv.org/abs/1905.10044",
    "google_fleurs": "https://arxiv.org/abs/2205.12446",
    "ibrahimhamamci_CT-RATE": "https://arxiv.org/abs/2403.17834",
    "openai_gsm8k": "https://arxiv.org/abs/2110.14168",
    "openai_openai_humaneval": "https://arxiv.org/abs/2107.03374",
    "rajpurkar_squad": "https://arxiv.org/abs/1606.05250",
    "truthfulqa_truthful_qa": "https://arxiv.org/abs/2109.07958",
    "Cnam-LMSSC_vibravox": "https://arxiv.org/abs/2407.11828",
    "EleutherAI_lambada_openai": "https://arxiv.org/abs/1606.06031",
    "IGNF_PASTIS-HD": "https://arxiv.org/abs/2107.07933",
    "Idavidrein_gpqa": "https://arxiv.org/abs/2311.12022",
    "Lin-Chen_MMStar": "https://arxiv.org/abs/2403.20330",
    "MohamedRashad_arabic-books": "https://arxiv.org/abs/2411.17835",
    "OpenGVLab_MVBench": "https://arxiv.org/abs/2311.17005",
    "allenai_math_qa": "https://arxiv.org/abs/1905.13319",
    "allenai_winogrande": "https://arxiv.org/abs/1907.10641",
    "callanwu_WebWalkerQA": "https://arxiv.org/abs/2501.07572",
    "ceval_ceval-exam": "https://arxiv.org/abs/2305.08322",
    "common-canvas_commoncatalog-cc-by": "https://arxiv.org/abs/2310.16825",
    "deepmind_code_contests": "https://arxiv.org/abs/2203.07814",
    "eriktks_conll2003": "https://arxiv.org/abs/cs/0306050",
    "google-research-datasets_mbpp": "https://arxiv.org/abs/2108.07732",
    "google_IFEval": "https://arxiv.org/abs/2311.07911",
    "hiyouga_geometry3k": "https://arxiv.org/abs/2105.04165",
    "lmms-lab_ChartQA": "https://arxiv.org/abs/2203.10244",
    "lmms-lab_DocVQA": "https://arxiv.org/abs/2007.00398",
    "lmms-lab_POPE": "https://arxiv.org/abs/2305.10355",
    "lmms-lab_ScienceQA": "https://arxiv.org/abs/2209.09513",
    "lmms-lab_Video-MME": "https://arxiv.org/abs/2405.21075",
    "lmms-lab_textvqa": "https://arxiv.org/abs/1904.08920",
    "locuslab_TOFU": "https://arxiv.org/abs/2401.06121",
    "lukaemon_bbh": "https://arxiv.org/abs/2210.09261",
    "m-a-p_SuperGPQA": "https://arxiv.org/abs/2502.14739",
    "princeton-nlp_SWE-bench": "https://arxiv.org/abs/2310.06770",
    "russellyq_VQA": "https://arxiv.org/abs/1610.01465",
    "tanganke_eurosat": "https://arxiv.org/abs/1709.00029",
    "tau_commonsense_qa": "https://arxiv.org/abs/1811.00937",
    "uoft-cs_cifar100": "https://www.cs.toronto.edu/~kriz/learning-features-2009-TR.pdf",
}

# ═══════════════════════════════════════════════════════════════════════
# Field definitions (from Croissant / Croissant-RAI spec)
# ═══════════════════════════════════════════════════════════════════════

FIELD_DEFINITIONS = [
    # General fields
    ("sc:name", "The name of the dataset.",
     "MMLU (Massive Multitask Language Understanding)"),
    ("sc:description", "A brief description of the dataset.",
     "A benchmark test to measure a text model's multitask accuracy across 57 tasks."),
    ("sc:url", "URL where the dataset can be accessed or downloaded.",
     "https://github.com/hendrycks/test"),
    ("sc:license", "The license under which the dataset is distributed.",
     "MIT License"),
    ("sc:creator", "The creator(s) or author(s) of the dataset.",
     "Dan Hendrycks, Collin Burns, Steven Basart, Andy Zou (UC Berkeley)"),
    ("sc:publisher", "The organization or entity that published the dataset.",
     "ICLR 2021"),
    ("sc:datePublished", "The date the dataset was first published or released.",
     "2021"),
    ("sc:inLanguage", "The language(s) of the dataset content.",
     "English"),
    ("cr:citeAs", "The recommended citation format for the dataset.",
     "@article{hendryckstest2021, title={Measuring Massive Multitask...}, year={2021}}"),
    ("cr:isLiveDataset", "Whether the dataset is actively updated (Yes/No).",
     "No"),
    # RAI fields
    ("rai:dataCollection", "Description of the data collection process: methodology, sources, tools used.",
     "Questions were manually collected by graduate and undergraduate students from freely available sources online, including practice exams for GRE and USMLE."),
    ("rai:dataCollectionType", "The type(s) of data collection. Choose from: Surveys, Secondary Data analysis, Physical data collection, Direct measurement, Document analysis, Manual Human Curator, Software Collection, Experiments, Web Scraping, Web API, Focus groups, Self-reporting, Customer feedback data, User-generated content data, Passive Data Collection, Others.",
     "Manual Human Curator"),
    ("rai:dataCollectionMissingData", "Description of how missing data was handled, if explicitly discussed.",
     "Images with fewer than 3 valid annotations were excluded from the final dataset."),
    ("rai:dataCollectionRawData", "Description of the raw/source data before processing.",
     "Questions from PDFs or websites where questions and answers are on separate pages."),
    ("rai:dataCollectionTimeframe", "The timeframe (start and end dates) of the data collection process.",
     "January 2019 to December 2020"),
    ("rai:dataImputationProtocol", "How missing or incomplete data values were imputed or filled.",
     "Missing demographic fields were imputed using k-nearest neighbors based on geographic region."),
    ("rai:dataManipulationProtocol", "Description of data manipulation procedures applied after preprocessing.",
     "Each source dataset was limited to 400 examples to ensure balanced representation."),
    ("rai:dataPreprocessingProtocol", "Steps to bring collected data to a state processable by ML models (e.g., filtering, cleaning, formatting).",
     "To encode mathematics expressions, we use LaTeX or symbols such as * and ^ for multiplication and exponentiation."),
    ("rai:dataAnnotationProtocol", "Description of how annotations (labels, ratings) were created or authored. Includes: annotation task, annotator instructions, workforce type, quality control.",
     "Crowdworkers on MTurk formulated questions about Wikipedia passages. Workers required 97% HIT acceptance rate. Each question received 3 independent answers."),
    ("rai:dataAnnotationPlatform", "Platform, tool, or library used to collect annotations by human annotators.",
     "Amazon Mechanical Turk"),
    ("rai:dataAnnotationAnalysis", "Analysis of annotation quality: uncertainty, disagreement, inter-annotator agreement.",
     "Inter-annotator agreement using Fleiss Kappa was 0.775, indicating substantial consistency."),
    ("rai:annotationsPerItem", "Number of human labels or annotations per data item.",
     "3 annotations per item"),
    ("rai:annotatorDemographics", "Demographic specifications about the annotators.",
     "Graduate and undergraduate students in STEM fields"),
    ("rai:machineAnnotationTools", "Software or ML tools used for data annotation (e.g., NER tools, automated labelers).",
     "EasyOCR for text detection; GPT-4 as answer extractor with 99.5% accuracy"),
    ("rai:dataReleaseMaintenancePlan", "Versioning information: updating timeframe, maintainers, deprecation policies.",
     "The dataset is maintained by the LMSYS team with quarterly updates. Version 2.0 released March 2024."),
    ("rai:personalSensitiveInformation", "Any sensitive human attributes collected (e.g., gender, socio-economic status, geography, age).",
     "The dataset contains face images with inferred age and gender attributes."),
    ("rai:dataSocialImpact", "Discussion of social implications of the dataset, if applicable.",
     "Models perform poorly on morality and law tasks, which is concerning for value alignment."),
    ("rai:dataBiases", "Description of known biases in the dataset, as explicitly discussed by the authors.",
     "Questions biased toward North American educational standards; English-only sources limit cross-lingual generalization."),
    ("rai:dataLimitations", "Known limitations of the dataset and non-recommended uses.",
     "The benchmark is text-only and does not incorporate multimodal information."),
    ("rai:dataUseCases", "Intended use cases for the dataset (e.g., Training, Testing, Validation, Fine Tuning).",
     "Testing and evaluation of pretrained language models on multitask understanding."),
]

# ═══════════════════════════════════════════════════════════════════════
# Styles
# ═══════════════════════════════════════════════════════════════════════

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LINK_FONT = Font(color="0563C1", underline="single", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _copy_cell_style(src, dst):
    """Copy style from src cell to dst cell."""
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)
    dst.number_format = src.number_format


# ═══════════════════════════════════════════════════════════════════════
# Update functions
# ═══════════════════════════════════════════════════════════════════════


def update_annotations_sheet(ws):
    """Update Annotations tab: insert Paper URL column, update rating dropdown."""
    max_row = ws.max_row

    # ── Step 1: Insert Paper URL column at position C (after Dataset ID) ──
    ws.insert_cols(3)

    # Set header for new column C
    ws.cell(1, 3, "Paper URL")
    ws.cell(1, 3).font = copy(ws.cell(1, 2).font) if ws.cell(1, 2).font else HEADER_FONT
    ws.cell(1, 3).fill = copy(ws.cell(1, 2).fill) if ws.cell(1, 2).fill else HEADER_FILL
    ws.cell(1, 3).alignment = copy(ws.cell(1, 2).alignment) if ws.cell(1, 2).alignment else HEADER_ALIGN
    ws.column_dimensions["C"].width = 40.0

    # Populate Paper URL for each row based on Dataset ID (column B)
    for r in range(2, max_row + 1):
        ds_id = ws.cell(r, 2).value  # Column B = Dataset ID
        if ds_id and ds_id in PAPER_URLS:
            url = PAPER_URLS[ds_id]
            ws.cell(r, 3, url)
            ws.cell(r, 3).font = LINK_FONT
            ws.cell(r, 3).hyperlink = url

    # ── Step 2: Update Rating column header (now column F after insert) ──
    # After inserting col C, old E→F
    ws.cell(1, 6, "Rating (1=Correct, 2=Partial, 3=Incorrect)")
    ws.column_dimensions["F"].width = 18.0

    # ── Step 3: Clear old data validations and re-add ──
    ws.data_validations.dataValidation.clear()

    # Rating dropdown (column F) — 1-3 with descriptive labels
    dv_rating = DataValidation(
        type="list",
        formula1='"1 - Correct,2 - Partially Correct,3 - Not Correct"',
        allow_blank=True,
    )
    dv_rating.error = "Please select 1, 2, or 3"
    dv_rating.errorTitle = "Invalid Rating"
    dv_rating.prompt = "1=Correct, 2=Partially Correct, 3=Not Correct"
    dv_rating.promptTitle = "Rating"
    for r in range(2, max_row + 1):
        dv_rating.add(ws.cell(r, 6))  # Column F
    ws.add_data_validation(dv_rating)

    # Failure Mode dropdown (column G)
    dv_failure = DataValidation(
        type="list",
        formula1='"Hallucination,Incomplete,Wrong Field,Format Mismatch,Overinterpreted,Other"',
        allow_blank=True,
    )
    for r in range(2, max_row + 1):
        dv_failure.add(ws.cell(r, 7))  # Column G
    ws.add_data_validation(dv_failure)

    # Confidence dropdown (column I)
    dv_confidence = DataValidation(
        type="list",
        formula1='"High,Medium,Low"',
        allow_blank=True,
    )
    for r in range(2, max_row + 1):
        dv_confidence.add(ws.cell(r, 9))  # Column I
    ws.add_data_validation(dv_confidence)


def update_instructions_sheet(ws):
    """Update Instructions tab to reflect 1-3 scale."""
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(r, 1)
        val = cell.value
        if val is None:
            continue
        v = str(val)

        # Update rating scale references
        if "Rating (1-5)" in v:
            cell.value = v.replace("Rating (1-5)", "Rating (1-3)")
        if "1 (Incorrect) to 5 (Correct)" in v:
            cell.value = cell.value.replace(
                "1 (Incorrect) to 5 (Correct)",
                "1 (Correct), 2 (Partially Correct), or 3 (Not Correct)"
            )
        if "select a rating from 1 (Correct)" in v:
            cell.value = cell.value.replace(
                "select a rating from 1 (Correct), 2 (Partially Correct), or 3 (Not Correct)",
                "select a rating: 1 (Correct), 2 (Partially Correct), or 3 (Not Correct)"
            )
        if "If Rating < 5:" in v:
            cell.value = cell.value.replace("If Rating < 5:", "If Rating > 1:")
        if "rate [NULL] extractions as 5 (Correct)" in v:
            cell.value = cell.value.replace(
                "rate [NULL] extractions as 5 (Correct)",
                "rate [NULL] extractions as 1 (Correct) — i.e., '1 - Correct'"
            )
        if "rate as 1 with 'Hallucination' failure mode" in v:
            cell.value = cell.value.replace(
                "rate as 1 with 'Hallucination' failure mode",
                "rate as 3 (Not Correct) with 'Hallucination' failure mode"
            )
        # RATING SCALE section — replace entirely
        elif v.strip() == "RATING SCALE:":
            cell.value = "RATING SCALE:"
            # Clear old rows 30-38 and write new scale
            # Row after "RATING SCALE:"
            row = r + 1
            new_scale = [
                "1 - Correct | Accurate and complete extraction.",
                "  Example: AI: 'The dataset was collected via Amazon Mechanical Turk' / Paper says the same.",
                "2 - Partially Correct | Captures some relevant information but incomplete or partially inaccurate.",
                "  Example: AI: 'Data was collected online' / Paper: 'Web scraping of Reddit posts from 2019-2021 using PRAW API'. Vague vs specific.",
                "3 - Not Correct | Completely wrong, irrelevant, hallucinated, or missing key information.",
                "  Example: AI: 'Dataset is updated quarterly with new samples' / Paper says nothing about updates.",
                None,  # blank row
                "NOTE: When a field genuinely has no info in the paper, [NULL] extractions are Correct (rate 1).",
                "When the AI extracted something but the paper says nothing, rate as 3 (Not Correct) with 'Hallucination'.",
            ]
            for i, line in enumerate(new_scale):
                ws.cell(row + i, 1).value = line

            # Clear any remaining old scale rows (the old 5-point had more lines)
            for clear_r in range(row + len(new_scale), row + 12):
                ws.cell(clear_r, 1).value = None

        # Update failure mode instruction
        elif "required when Rating < 5" in v:
            cell.value = v.replace("required when Rating < 5", "required when Rating > 1")

        # Update about hallucination section
        elif "should be rated 5 (Correct)" in v:
            cell.value = v.replace("should be rated 5 (Correct)", "should be rated 1 (Correct)")
        elif "should be rated 1 (Incorrect)" in v:
            cell.value = v.replace("should be rated 1 (Incorrect)", "should be rated 3 (Not Correct)")


def create_field_definitions_sheet(wb):
    """Add a 'Field Definitions' tab with all 30 field definitions."""
    if "Field Definitions" in wb.sheetnames:
        del wb["Field Definitions"]

    ws = wb.create_sheet("Field Definitions")

    # Headers
    headers = ["Field Name", "Definition (Croissant Spec)", "Example of a Good Value"]
    col_widths = [35, 80, 70]

    for c, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(1, c, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(c)].width = width

    # Data rows
    for i, (field, definition, example) in enumerate(FIELD_DEFINITIONS, 2):
        ws.cell(i, 1, field).font = Font(bold=True, size=10)
        ws.cell(i, 1).border = THIN_BORDER
        ws.cell(i, 2, definition).alignment = Alignment(wrap_text=True)
        ws.cell(i, 2).border = THIN_BORDER
        ws.cell(i, 3, example).alignment = Alignment(wrap_text=True)
        ws.cell(i, 3).font = Font(italic=True, size=10, color="555555")
        ws.cell(i, 3).border = THIN_BORDER

    # Freeze header row
    ws.freeze_panes = "A2"

    # Add a separator between General and RAI fields
    # General fields are rows 2-11 (10 fields), RAI starts at row 12
    for c in range(1, 4):
        cell = ws.cell(12, c)
        cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")


def update_workbook(filepath: Path):
    """Apply all changes to a single workbook."""
    wb = openpyxl.load_workbook(filepath)

    # 1. Update Annotations sheet
    if "Annotations" in wb.sheetnames:
        update_annotations_sheet(wb["Annotations"])

    # 2. Update Instructions sheet
    if "Instructions" in wb.sheetnames:
        update_instructions_sheet(wb["Instructions"])

    # 3. Add Field Definitions tab
    create_field_definitions_sheet(wb)

    # Reorder sheets: Instructions, Annotations, Field Definitions
    desired_order = []
    for name in ["Instructions", "Annotations", "Field Definitions"]:
        if name in wb.sheetnames:
            desired_order.append(wb.sheetnames.index(name))
    if len(desired_order) == 3:
        wb.move_sheet("Field Definitions", offset=-(len(wb.sheetnames) - 3))

    wb.save(filepath)


# ═══════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════

def main():
    xlsx_files = sorted(PHASE2_DIR.glob("*.xlsx"))
    print(f"Found {len(xlsx_files)} sheets in {PHASE2_DIR}\n")

    for i, filepath in enumerate(xlsx_files, 1):
        name = filepath.stem
        print(f"  [{i:2d}/{len(xlsx_files)}] Updating {name}...")
        try:
            update_workbook(filepath)
        except Exception as e:
            print(f"    ERROR: {e}")

    print(f"\nDone. Updated {len(xlsx_files)} sheets.")
    print(f"\nVerify with: python3 -c \"import openpyxl; wb=openpyxl.load_workbook('{xlsx_files[0]}'); print(wb.sheetnames)\"")


if __name__ == "__main__":
    main()
