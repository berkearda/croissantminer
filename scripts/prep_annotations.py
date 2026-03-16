#!/usr/bin/env python3
"""
Prepare annotation data for Google Sheets-based validation workflow.

This script:
1. Loads all extracted metadata from data/processed/*/full_pdf_metadata_result.json
2. Assigns fields to annotators (3 fields per annotator)
3. Generates formatted Excel files ready for Google Sheets import
4. Creates annotator assignment summary

Usage:
    python scripts/prep_annotations.py
"""

import json
import os
from pathlib import Path
from datetime import datetime
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# Configuration
PROCESSED_DIR = Path("data/processed")
OUTPUT_DIR = Path("data/annotations")
PAPER_LINKS_FILE = Path("data/paper_links.json")
FIELDS_PER_ANNOTATOR = 3

# All 30 fields in the schema
ALL_FIELDS = {
    "general": [
        "name", "description", "url", "license", "creator",
        "publisher", "datePublished", "inLanguage", "citeAs", "isLiveDataset"
    ],
    "rai": [
        "rai:dataCollection", "rai:dataCollectionType", "rai:dataCollectionMissingData",
        "rai:dataCollectionRawData", "rai:dataCollectionTimeframe", "rai:dataImputationProtocol",
        "rai:dataManipulationProtocol", "rai:dataPreprocessingProtocol", "rai:dataAnnotationProtocol",
        "rai:dataAnnotationPlatform", "rai:dataAnnotationAnalysis", "rai:annotationsPerItem",
        "rai:annotatorDemographics", "rai:machineAnnotationTools", "rai:dataReleaseMaintenancePlan",
        "rai:personalSensitiveInformation", "rai:dataSocialImpact", "rai:dataBiases",
        "rai:dataLimitations", "rai:dataUseCases"
    ]
}

# Annotator assignments (customize names as needed)
ANNOTATOR_ASSIGNMENTS = {
    "Annotator_01": ["name", "description", "url"],
    "Annotator_02": ["license", "creator", "publisher"],
    "Annotator_03": ["datePublished", "inLanguage", "citeAs"],
    "Annotator_04": ["isLiveDataset", "rai:dataCollection", "rai:dataCollectionType"],
    "Annotator_05": ["rai:dataCollectionMissingData", "rai:dataCollectionRawData", "rai:dataCollectionTimeframe"],
    "Annotator_06": ["rai:dataImputationProtocol", "rai:dataManipulationProtocol", "rai:dataPreprocessingProtocol"],
    "Annotator_07": ["rai:dataAnnotationProtocol", "rai:dataAnnotationPlatform", "rai:dataAnnotationAnalysis"],
    "Annotator_08": ["rai:annotationsPerItem", "rai:annotatorDemographics", "rai:machineAnnotationTools"],
    "Annotator_09": ["rai:dataReleaseMaintenancePlan", "rai:personalSensitiveInformation", "rai:dataSocialImpact"],
    "Annotator_10": ["rai:dataBiases", "rai:dataLimitations", "rai:dataUseCases"],
}

# Color scheme
COLORS = {
    "header_bg": "1F4E79",      # Dark blue
    "header_font": "FFFFFF",    # White
    "row_alt_1": "FFFFFF",      # White
    "row_alt_2": "D6EAF8",      # Light blue
    "verdict_bg": "FFF9C4",     # Light yellow
    "correction_bg": "FFE0B2",  # Light orange
    "confidence_bg": "C8E6C9",  # Light green
    "completed_true": "C8E6C9", # Light green
    "completed_false": "FFCDD2", # Light red
}

# Column configuration
COLUMNS = [
    {"name": "#", "width": 6},
    {"name": "Dataset ID", "width": 28},
    {"name": "Field Name", "width": 32},
    {"name": "Extracted Value", "width": 80},
    {"name": "Verdict", "width": 12},
    {"name": "Corrected Value", "width": 50},
    {"name": "Confidence", "width": 12},
    {"name": "Notes", "width": 35},
]


def load_paper_links():
    """Load paper links mapping from JSON file."""
    if PAPER_LINKS_FILE.exists():
        with open(PAPER_LINKS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def load_all_extractions():
    """Load all extracted metadata from processed directories."""
    extractions = {}

    for dataset_dir in sorted(PROCESSED_DIR.iterdir()):
        if not dataset_dir.is_dir():
            continue

        # Try both possible filenames
        for filename in ["full_pdf_metadata_result.json", "croissant_metadata.json"]:
            result_file = dataset_dir / filename
            if result_file.exists():
                try:
                    with open(result_file, 'r', encoding='utf-8') as f:
                        data = json.load(f)

                    # Only include if it has RAI fields (30-field schema)
                    if "rai:dataCollection" in data or "rai:dataUseCases" in data:
                        extractions[dataset_dir.name] = data
                        break
                except json.JSONDecodeError:
                    print(f"  Warning: Invalid JSON in {result_file}")
                except Exception as e:
                    print(f"  Warning: Error loading {result_file}: {e}")

    return extractions


def format_value_for_display(value):
    """Format a value for display in Excel."""
    if value is None:
        return "[NULL]"
    elif isinstance(value, (list, dict)):
        # Convert complex types to readable JSON
        formatted = json.dumps(value, indent=2, ensure_ascii=False)
        # Truncate if too long
        if len(formatted) > 3000:
            formatted = formatted[:2997] + "..."
        return formatted
    elif isinstance(value, str):
        if not value.strip():
            return "[EMPTY]"
        # Truncate very long strings
        if len(value) > 3000:
            return value[:2997] + "..."
        return value
    else:
        return str(value)


def create_instructions_sheet(wb, annotator_name, fields, extractions, paper_links):
    """Create an instructions sheet as the first tab."""
    ws = wb.active
    ws.title = "Instructions"

    # Title styling
    title_font = Font(name='Arial', size=16, bold=True, color=COLORS["header_bg"])
    header_font = Font(name='Arial', size=12, bold=True)
    normal_font = Font(name='Arial', size=11)
    link_font = Font(name='Arial', size=10, color="0563C1", underline="single")
    no_paper_font = Font(name='Arial', size=10, italic=True, color="888888")

    # Title
    ws['A1'] = f"Annotation Instructions - {annotator_name}"
    ws['A1'].font = title_font
    ws.merge_cells('A1:D1')

    # Overview section
    ws['A3'] = "YOUR ASSIGNED FIELDS:"
    ws['A3'].font = header_font

    row = 4
    for i, field in enumerate(fields, 1):
        ws[f'A{row}'] = f"  {i}. {field}"
        ws[f'A{row}'].font = normal_font
        row += 1

    row += 1
    ws[f'A{row}'] = "HOW TO ANNOTATE:"
    ws[f'A{row}'].font = header_font
    row += 1

    instructions = [
        "1. Go to the 'Annotations' sheet (tab at bottom)",
        "2. For each row, review the 'Extracted Value' column",
        "3. Compare with the original paper (use Paper Links below)",
        "4. In 'Verdict' column: Select TRUE if correct, FALSE if incorrect",
        "5. If FALSE: Type the correct value in 'Corrected Value' column",
        "6. In 'Confidence' column: Select High/Medium/Low",
        "7. Add any notes in the 'Notes' column (optional)",
    ]

    for instruction in instructions:
        ws[f'A{row}'] = instruction
        ws[f'A{row}'].font = normal_font
        row += 1

    row += 1
    ws[f'A{row}'] = "COLOR LEGEND:"
    ws[f'A{row}'].font = header_font
    row += 1

    legend = [
        ("Yellow cells", "Where to enter your Verdict (TRUE/FALSE)"),
        ("Orange cells", "Where to enter corrections (if FALSE)"),
        ("Green cells", "Where to enter your confidence (High/Medium/Low)"),
        ("Blue/white rows", "Different datasets (alternating)"),
    ]

    for label, desc in legend:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(name='Arial', size=11, bold=True)
        ws[f'B{row}'] = desc
        ws[f'B{row}'].font = normal_font
        row += 1

    # Paper links section
    row += 2
    ws[f'A{row}'] = "PAPER LINKS:"
    ws[f'A{row}'].font = header_font
    row += 1

    ws[f'A{row}'] = "Dataset ID"
    ws[f'A{row}'].font = Font(name='Arial', size=11, bold=True)
    ws[f'B{row}'] = "Paper URL"
    ws[f'B{row}'].font = Font(name='Arial', size=11, bold=True)
    row += 1

    for dataset_id in sorted(extractions.keys()):
        # Get paper URL from the paper_links mapping
        paper_url = paper_links.get(dataset_id, None)

        ws[f'A{row}'] = dataset_id
        ws[f'A{row}'].font = normal_font

        if paper_url:
            ws[f'B{row}'] = paper_url
            ws[f'B{row}'].font = link_font
            ws[f'B{row}'].hyperlink = paper_url
        else:
            ws[f'B{row}'] = "(no paper)"
            ws[f'B{row}'].font = no_paper_font

        row += 1

    # Set column widths
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 70

    return ws


def create_annotation_sheet(wb, annotator_name, fields, extractions):
    """Create the main annotation sheet with formatting."""
    ws = wb.create_sheet("Annotations")

    # Define styles
    header_fill = PatternFill(start_color=COLORS["header_bg"], end_color=COLORS["header_bg"], fill_type="solid")
    header_font = Font(name='Arial', size=11, bold=True, color=COLORS["header_font"])
    normal_font = Font(name='Arial', size=10)

    verdict_fill = PatternFill(start_color=COLORS["verdict_bg"], end_color=COLORS["verdict_bg"], fill_type="solid")
    correction_fill = PatternFill(start_color=COLORS["correction_bg"], end_color=COLORS["correction_bg"], fill_type="solid")
    confidence_fill = PatternFill(start_color=COLORS["confidence_bg"], end_color=COLORS["confidence_bg"], fill_type="solid")

    row_fill_1 = PatternFill(start_color=COLORS["row_alt_1"], end_color=COLORS["row_alt_1"], fill_type="solid")
    row_fill_2 = PatternFill(start_color=COLORS["row_alt_2"], end_color=COLORS["row_alt_2"], fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    center_alignment = Alignment(horizontal='center', vertical='center')

    # Write header row
    for col_idx, col_info in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_info["name"])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_info["width"]

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Add data validation for verdict (dropdown)
    verdict_dv = DataValidation(
        type="list",
        formula1='"TRUE,FALSE"',
        allow_blank=True
    )
    verdict_dv.error = "Please select TRUE or FALSE"
    verdict_dv.errorTitle = "Invalid Input"
    verdict_dv.prompt = "Select TRUE if correct, FALSE if incorrect"
    verdict_dv.promptTitle = "Verdict"
    ws.add_data_validation(verdict_dv)

    # Add data validation for confidence (dropdown)
    confidence_dv = DataValidation(
        type="list",
        formula1='"High,Medium,Low"',
        allow_blank=True
    )
    confidence_dv.error = "Please select High, Medium, or Low"
    confidence_dv.errorTitle = "Invalid Input"
    confidence_dv.prompt = "High=Very confident, Medium=Somewhat confident, Low=Uncertain"
    confidence_dv.promptTitle = "Confidence"
    ws.add_data_validation(confidence_dv)

    # Write data rows
    row_num = 2
    dataset_list = sorted(extractions.keys())

    for dataset_idx, dataset_id in enumerate(dataset_list):
        data = extractions[dataset_id]
        # Alternate colors by dataset
        row_fill = row_fill_2 if dataset_idx % 2 == 0 else row_fill_1

        for field in fields:
            extracted_value = data.get(field)
            formatted_value = format_value_for_display(extracted_value)

            # Row number
            cell = ws.cell(row=row_num, column=1, value=row_num - 1)
            cell.font = normal_font
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = center_alignment

            # Dataset ID
            cell = ws.cell(row=row_num, column=2, value=dataset_id)
            cell.font = normal_font
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top')

            # Field name
            cell = ws.cell(row=row_num, column=3, value=field)
            cell.font = Font(name='Arial', size=10, bold=True)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top')

            # Extracted value
            cell = ws.cell(row=row_num, column=4, value=formatted_value)
            cell.font = normal_font
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = wrap_alignment

            # Verdict (input cell - yellow)
            cell = ws.cell(row=row_num, column=5, value="")
            cell.fill = verdict_fill
            cell.border = thin_border
            cell.alignment = center_alignment
            verdict_dv.add(cell)

            # Corrected value (input cell - orange)
            cell = ws.cell(row=row_num, column=6, value="")
            cell.fill = correction_fill
            cell.border = thin_border
            cell.alignment = wrap_alignment

            # Confidence (input cell - green)
            cell = ws.cell(row=row_num, column=7, value="")
            cell.fill = confidence_fill
            cell.border = thin_border
            cell.alignment = center_alignment
            confidence_dv.add(cell)

            # Notes
            cell = ws.cell(row=row_num, column=8, value="")
            cell.font = normal_font
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = wrap_alignment

            row_num += 1

    # Set row height for better readability
    for row in range(2, row_num):
        ws.row_dimensions[row].height = 45

    return row_num - 2  # Return number of data rows


def create_annotator_xlsx(annotator_name, fields, extractions, output_dir, paper_links):
    """Create a formatted Excel file for a specific annotator."""
    output_file = output_dir / f"{annotator_name}.xlsx"

    wb = Workbook()

    # Create instructions sheet (first tab)
    create_instructions_sheet(wb, annotator_name, fields, extractions, paper_links)

    # Create annotation sheet (second tab)
    row_count = create_annotation_sheet(wb, annotator_name, fields, extractions)

    # Save workbook
    wb.save(output_file)

    return row_count


def create_assignment_summary(annotator_assignments, extractions, output_dir):
    """Create a summary of annotator assignments."""
    summary_file = output_dir / "annotator_assignments.md"

    num_datasets = len(extractions)

    with open(summary_file, 'w', encoding='utf-8') as f:
        f.write("# Annotation Assignment Summary\n\n")
        f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        f.write(f"## Overview\n\n")
        f.write(f"- **Total Datasets**: {num_datasets}\n")
        f.write(f"- **Total Fields**: 30 (10 General + 20 RAI)\n")
        f.write(f"- **Total Annotations**: {num_datasets * 30:,}\n")
        f.write(f"- **Annotators**: {len(annotator_assignments)}\n")
        f.write(f"- **Fields per Annotator**: {FIELDS_PER_ANNOTATOR}\n")
        f.write(f"- **Annotations per Annotator**: ~{num_datasets * FIELDS_PER_ANNOTATOR:,}\n\n")

        f.write("## Annotator Assignments\n\n")
        f.write("| Annotator | Assigned Fields | # Annotations |\n")
        f.write("|-----------|-----------------|---------------|\n")

        for annotator, fields in annotator_assignments.items():
            fields_str = ", ".join(f"`{f}`" for f in fields)
            num_annotations = num_datasets * len(fields)
            f.write(f"| {annotator} | {fields_str} | {num_annotations} |\n")

        f.write("\n## How to Use\n\n")
        f.write("1. Share the `.xlsx` file with each annotator\n")
        f.write("2. They can open it in Google Sheets or Excel\n")
        f.write("3. First tab has instructions, second tab has annotations\n")
        f.write("4. Collect completed files and run `analyze_annotations.py`\n")

    return summary_file


def main():
    """Main function to prepare annotation files."""
    print("=" * 60)
    print("CroissantMiner - Annotation Preparation (Excel Format)")
    print("=" * 60)

    # Create output directory
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\nOutput directory: {OUTPUT_DIR}")

    # Load extractions
    print("\nLoading extracted metadata...")
    extractions = load_all_extractions()
    print(f"  Loaded {len(extractions)} datasets with 30-field schema")

    if not extractions:
        print("ERROR: No valid extractions found!")
        return

    # Load paper links
    print("\nLoading paper links...")
    paper_links = load_paper_links()
    print(f"  Loaded {len(paper_links)} paper URLs")

    # Create per-annotator Excel files
    print("\nCreating formatted Excel files...")
    for annotator, fields in ANNOTATOR_ASSIGNMENTS.items():
        row_count = create_annotator_xlsx(annotator, fields, extractions, OUTPUT_DIR, paper_links)
        print(f"  {annotator}: {len(fields)} fields × {len(extractions)} datasets = {row_count} annotations")

    # Create assignment summary
    print("\nCreating assignment summary...")
    summary_file = create_assignment_summary(ANNOTATOR_ASSIGNMENTS, extractions, OUTPUT_DIR)
    print(f"  Created {summary_file}")

    # Print final summary
    total_annotations = len(extractions) * 30
    print("\n" + "=" * 60)
    print("PREPARATION COMPLETE")
    print("=" * 60)
    print(f"\nFiles created in {OUTPUT_DIR}/:")
    print(f"  - Annotator_01.xlsx through Annotator_10.xlsx")
    print(f"  - annotator_assignments.md (instructions)")
    print(f"\nTotal annotations to collect: {total_annotations:,}")
    print(f"Annotations per annotator: ~{total_annotations // 10:,}")
    print("\nFeatures included:")
    print("  ✓ Alternating row colors by dataset")
    print("  ✓ Color-coded input columns (verdict/correction/confidence)")
    print("  ✓ Dropdown menus for Verdict and Confidence")
    print("  ✓ Frozen header row")
    print("  ✓ Instructions sheet as first tab")
    print("\nNext steps:")
    print("  1. Share .xlsx files with annotators")
    print("  2. They can open in Google Sheets or Excel")
    print("  3. Collect completed files and run analyze_annotations.py")


if __name__ == "__main__":
    main()
